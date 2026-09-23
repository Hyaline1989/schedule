from __future__ import annotations

import os
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from schedule.bitrix import BitrixClient
from schedule.config import load_env_file
from schedule.sheet_objects import load_sheet_object_mappings


OBJECT_FIELD = "UF_CRM_642A972B79B02"
WORKED_SHIFTS_FIELD = "UF_CRM_WORKED_SHIFTS"


def main() -> None:
    load_env_file()
    category_id = int(os.environ.get("BITRIX_CATEGORY_ID", "0"))
    start_date = date.fromisoformat(
        os.environ.get("REPORT_START_DATE")
        or os.environ.get("SCHEDULE_START_DATE", "2026-08-01")
    )
    end_date = date.fromisoformat(
        os.environ.get("REPORT_END_DATE")
        or os.environ.get("SCHEDULE_END_DATE")
        or date.today().isoformat()
    )
    worked_shifts_field = os.environ.get(
        "BITRIX_WORKED_SHIFTS_FIELD",
        WORKED_SHIFTS_FIELD,
    )
    bitrix = BitrixClient.from_environment()

    deals = bitrix._list_deals(
        {
            "CATEGORY_ID": category_id,
            ">=DATE_CREATE": start_date.isoformat(),
            "<=DATE_CREATE": f"{end_date.isoformat()}T23:59:59",
        }
    )
    object_names = _object_names()

    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    skipped_without_object = 0
    skipped_without_shifts = 0
    skipped_non_numeric = 0

    for deal in deals:
        shifts = _integer_or_none(deal.get(worked_shifts_field))
        if shifts is None:
            if deal.get(worked_shifts_field) in (None, ""):
                skipped_without_shifts += 1
            else:
                skipped_non_numeric += 1
            continue
        object_id = _object_id(deal.get(OBJECT_FIELD))
        if not object_id:
            skipped_without_object += 1
            continue
        groups[object_id].append({**deal, "_worked_shifts": shifts})

    output_dir = Path("outputs")
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / "worked_shifts_by_object.xlsx"
    _write_workbook(
        output_path,
        groups=groups,
        object_names=object_names,
        start_date=start_date,
        end_date=end_date,
        total_deals=len(deals),
        skipped_without_object=skipped_without_object,
        skipped_without_shifts=skipped_without_shifts,
        skipped_non_numeric=skipped_non_numeric,
    )
    print(f"created: {output_path}")
    print(f"objects: {len(groups)}")
    print(f"deals with worked shifts and object: {sum(len(items) for items in groups.values())}")
    print(f"deals in period: {len(deals)}")


def _object_names() -> dict[str, str]:
    result: dict[str, str] = {}
    for mapping in load_sheet_object_mappings().values():
        result[str(mapping.company_id)] = mapping.object_name or f"ID {mapping.company_id}"
    return result


def _write_workbook(
    path: Path,
    *,
    groups: dict[str, list[dict[str, Any]]],
    object_names: dict[str, str],
    start_date: date,
    end_date: date,
    total_deals: int,
    skipped_without_object: int,
    skipped_without_shifts: int,
    skipped_non_numeric: int,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "По объектам"

    ws.append(
        [
            "ID объекта",
            "Объект",
            "Количество сделок с заполненным количеством смен",
            "Среднее количество смен",
            "Минимум смен",
            "Максимум смен",
            "Сумма смен",
        ]
    )

    for object_id, items in sorted(
        groups.items(),
        key=lambda pair: (object_names.get(pair[0]) or "", int(pair[0]) if pair[0].isdigit() else pair[0]),
    ):
        values = [int(item["_worked_shifts"]) for item in items]
        ws.append(
            [
                object_id,
                object_names.get(object_id, ""),
                len(values),
                round(mean(values), 2),
                min(values),
                max(values),
                sum(values),
            ]
        )

    ws_summary = wb.create_sheet("Сводка")
    ws_summary.append(["Показатель", "Значение"])
    ws_summary.append(["Дата формирования", datetime.now().strftime("%d.%m.%Y %H:%M:%S")])
    ws_summary.append(["Период сделок", f"{start_date.isoformat()} — {end_date.isoformat()}"])
    ws_summary.append(["Сделок в периоде", total_deals])
    ws_summary.append(
        [
            "Сделок с заполненным количеством смен и объектом",
            sum(len(items) for items in groups.values()),
        ]
    )
    ws_summary.append(["Объектов в отчёте", len(groups)])
    ws_summary.append(["Пропущено: не заполнено количество смен", skipped_without_shifts])
    ws_summary.append(["Пропущено: количество смен не число", skipped_non_numeric])
    ws_summary.append(["Пропущено: не заполнен объект", skipped_without_object])

    ws_details = wb.create_sheet("Сделки")
    ws_details.append(
        [
            "ID объекта",
            "Объект",
            "ID сделки",
            "Ссылка на сделку",
            "ФИО / название сделки",
            "Стадия",
            "Количество смен",
        ]
    )
    for object_id, items in sorted(groups.items()):
        for deal in sorted(items, key=lambda item: int(str(item.get("ID") or 0))):
            deal_id = str(deal.get("ID") or "")
            ws_details.append(
                [
                    object_id,
                    object_names.get(object_id, ""),
                    deal_id,
                    f"https://crm.complect.org/crm/deal/details/{deal_id}/" if deal_id else "",
                    deal.get("TITLE") or "",
                    deal.get("STAGE_ID") or "",
                    deal["_worked_shifts"],
                ]
            )

    for sheet in wb.worksheets:
        _format_sheet(sheet)

    wb.save(path)


def _format_sheet(ws: Any) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, ws.max_column + 1):
        letter = get_column_letter(column)
        width = 16
        if column in (2, 5):
            width = 34
        elif column == 4:
            width = 42
        elif column in (3,):
            width = 22
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _object_id(value: Any) -> str:
    if value in (None, "", False):
        return ""
    if isinstance(value, list):
        return str(value[0]).strip() if value else ""
    return str(value).strip()


def _integer_or_none(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(",", ".").strip()))
    except (TypeError, ValueError):
        return None


if __name__ == "__main__":
    main()
