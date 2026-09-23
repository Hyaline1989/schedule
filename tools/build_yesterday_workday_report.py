from __future__ import annotations

import os
import shutil
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from schedule.browser_sheets import BrowserGoogleSheetsReader
from schedule.bitrix import BitrixClient
from schedule.config import load_env_file
from schedule.models import normalize_employee_name
from schedule.sheet_objects import load_sheet_object_mappings
from schedule.sources import spreadsheet_ids
from schedule.xlsx_schedules import parse_month_sheets


def main() -> None:
    load_env_file()
    report_date = date.fromisoformat(
        os.environ.get("REPORT_DATE", (date.today() - timedelta(days=1)).isoformat())
    )
    runtime_dir = Path("runtime")
    runtime_dir.mkdir(exist_ok=True)
    output_dir = Path("outputs")
    output_dir.mkdir(exist_ok=True)

    mappings = load_sheet_object_mappings()
    object_names = _object_names(mappings)
    reader = BrowserGoogleSheetsReader(
        os.environ.get("GOOGLE_CDP_URL", "http://127.0.0.1:9222"),
        timeout_ms=int(os.environ.get("SCHEDULE_GOOGLE_TIMEOUT_MS", "180000")),
    )

    summary: dict[str, dict[str, Any]] = {}
    details = []
    warnings = []
    ids = spreadsheet_ids()
    use_cached = os.environ.get("SCHEDULE_USE_CACHED_XLSX", "").strip() == "1"

    for index, spreadsheet_id in enumerate(ids, start=1):
        mapping = mappings.get(spreadsheet_id)
        object_id = mapping.company_id if mapping else ""
        object_name = object_names.get(str(object_id), f"ID {object_id}" if object_id else "")
        target = runtime_dir / f"{spreadsheet_id}.xlsx"
        try:
            if use_cached and target.exists():
                print(f"[{index}/{len(ids)}] cached {spreadsheet_id}", flush=True)
            else:
                print(f"[{index}/{len(ids)}] export {spreadsheet_id}", flush=True)
                reader.export_xlsx(spreadsheet_id, target)
            parsed = parse_month_sheets(
                target,
                source_id=spreadsheet_id,
                start_date=report_date,
                end_date=report_date,
            )
        except Exception as error:
            warnings.append(
                {
                    "spreadsheet_id": spreadsheet_id,
                    "object_id": object_id,
                    "object_name": object_name,
                    "error": str(error),
                }
            )
            print(f"[{index}/{len(ids)}] failed {spreadsheet_id}: {error}", flush=True)
            continue

        unique_by_name = {}
        for shift in parsed.shifts:
            key = normalize_employee_name(shift.employee_name)
            unique_by_name.setdefault(key, shift)

        summary[spreadsheet_id] = {
            "spreadsheet_id": spreadsheet_id,
            "object_id": object_id,
            "object_name": object_name,
            "people_count": len(unique_by_name),
            "raw_shift_rows": len(parsed.shifts),
        }

        for shift in sorted(
            unique_by_name.values(),
            key=lambda item: normalize_employee_name(item.employee_name),
        ):
            details.append(
                {
                    "object_id": object_id,
                    "object_name": object_name,
                    "employee": shift.employee_name,
                    "work_date": shift.work_date.isoformat(),
                    "cell_value": str(shift.cell_value),
                    "section": shift.section.value,
                    "source": shift.source_id or "",
                    "spreadsheet_id": spreadsheet_id,
                }
            )

        print(
            f"[{index}/{len(ids)}] people={len(unique_by_name)} shiftRows={len(parsed.shifts)}",
            flush=True,
        )

    output_path = output_dir / f"workday_people_{report_date.isoformat()}.xlsx"
    _write_workbook(output_path, report_date, summary.values(), details, warnings)
    public_dir = Path(
        os.environ.get(
            "SCHEDULE_PUBLIC_REPORTS_DIR",
            "/var/www/maxbot_compl_usr/data/www/maxbot.complect.org/schedule-reports",
        )
    )
    try:
        public_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(output_path, public_dir / output_path.name)
        print(
            "url: "
            f"https://maxbot.complect.org/schedule-reports/{output_path.name}"
        )
    except Exception as error:
        print(f"public copy failed: {error}")
    print(f"created: {output_path}")
    print(f"date: {report_date.isoformat()}")
    print(f"objects: {len(summary)}")
    print(f"people total by object sum: {sum(item['people_count'] for item in summary.values())}")
    print(f"warnings: {len(warnings)}")


def _object_names(mappings: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for mapping in mappings.values():
        result[str(mapping.company_id)] = mapping.object_name or f"ID {mapping.company_id}"
    try:
        bitrix = BitrixClient.from_environment()
    except Exception:
        return result
    for company_id, current_name in list(result.items()):
        if current_name != f"ID {company_id}":
            continue
        try:
            company = bitrix.get_company(company_id)
        except Exception:
            continue
        title = str(company.get("TITLE") or "").strip()
        if title:
            result[company_id] = title
    return result


def _write_workbook(
    path: Path,
    report_date: date,
    summary_rows: Any,
    details: list[dict[str, Any]],
    warnings: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "По объектам"
    ws.append(
        [
            "Дата",
            "ID объекта",
            "Объект",
            "Количество людей с рабочим днём",
            "Количество строк смен",
            "Ссылка на Google-таблицу",
        ]
    )
    for row in sorted(
        summary_rows,
        key=lambda item: (str(item.get("object_name") or ""), str(item.get("object_id") or "")),
    ):
        spreadsheet_id = row["spreadsheet_id"]
        ws.append(
            [
                report_date.isoformat(),
                row["object_id"],
                row["object_name"],
                row["people_count"],
                row["raw_shift_rows"],
                f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit",
            ]
        )

    ws_details = wb.create_sheet("Люди")
    ws_details.append(
        [
            "ID объекта",
            "Объект",
            "ФИО",
            "Дата",
            "Значение ячейки",
            "Секция",
            "Источник",
            "Ссылка на Google-таблицу",
        ]
    )
    for row in details:
        ws_details.append(
            [
                row["object_id"],
                row["object_name"],
                row["employee"],
                row["work_date"],
                row["cell_value"],
                row["section"],
                row["source"],
                f"https://docs.google.com/spreadsheets/d/{row['spreadsheet_id']}/edit",
            ]
        )

    ws_warnings = wb.create_sheet("Ошибки")
    ws_warnings.append(["ID объекта", "Объект", "Google-таблица", "Ошибка"])
    for warning in warnings:
        spreadsheet_id = warning["spreadsheet_id"]
        ws_warnings.append(
            [
                warning["object_id"],
                warning["object_name"],
                f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit",
                warning["error"],
            ]
        )

    for worksheet in wb.worksheets:
        _format_sheet(worksheet)
    wb.save(path)


def _format_sheet(worksheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column)
        width = 18
        if column in (3, 6, 7, 8):
            width = 34
        if column == 4:
            width = 24
        worksheet.column_dimensions[letter].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


if __name__ == "__main__":
    main()
