from __future__ import annotations

import json
import os
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PORTAL_URL = os.environ.get("BITRIX_PORTAL_URL", "https://crm.complect.org").rstrip("/")


def main() -> None:
    end_date = date.fromisoformat(os.environ.get("SCHEDULE_END_DATE", date.today().isoformat()))
    start_date = date.fromisoformat(os.environ.get("SCHEDULE_START_DATE", "2026-08-01"))
    runtime_dir = Path("runtime")
    output_dir = Path("outputs")
    output_dir.mkdir(exist_ok=True)

    report = json.loads((runtime_dir / "stage-mismatch-report.json").read_text(encoding="utf-8"))
    stage_path = output_dir / f"stage_mismatches_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"
    not_found_path = output_dir / f"not_found_in_bitrix_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"

    _write_stage_mismatches(report.get("rows", []), stage_path)
    _write_not_found(
        [row for row in report.get("unresolved", []) if row.get("reason") == "not_found"],
        not_found_path,
    )
    print(
        json.dumps(
            {
                "stageMismatches": len(report.get("rows", [])),
                "stageMismatchesXlsx": str(stage_path),
                "notFoundInBitrix": sum(
                    1 for row in report.get("unresolved", []) if row.get("reason") == "not_found"
                ),
                "notFoundInBitrixXlsx": str(not_found_path),
            },
            ensure_ascii=False,
        )
    )


def _write_stage_mismatches(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "stages"
    headers = [
        "ФИО из табеля",
        "ФИО в Bitrix",
        "Сделка",
        "ID объекта",
        "Сейчас",
        "Должно быть",
        "Смен",
        "Последняя смена",
        "Дата заселения",
        "Маркер",
        "Дата маркера",
        "Комментарий",
        "Таблица",
    ]
    worksheet.append(headers)
    for row in rows:
        deal_id = str(row.get("dealId") or "")
        sources = _source_urls(row.get("sources", []))
        worksheet.append(
            [
                row.get("employee"),
                row.get("matchedName"),
                f"{PORTAL_URL}/crm/deal/details/{deal_id}/" if deal_id else "",
                row.get("companyId"),
                row.get("currentStage"),
                row.get("desiredStage"),
                row.get("shiftCount"),
                row.get("lastShiftDate"),
                row.get("settlementDate"),
                row.get("statusMarker"),
                row.get("statusMarkerDate"),
                row.get("statusComment"),
                "\n".join(sources),
            ]
        )
        row_index = worksheet.max_row
        _hyperlink_cell(worksheet.cell(row_index, 3))
        _hyperlink_sources_cell(worksheet.cell(row_index, 13), sources)
    _format_table(worksheet, widths=[34, 34, 48, 12, 24, 28, 10, 14, 14, 12, 14, 34, 82])
    workbook.save(path)


def _write_not_found(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "not_found"
    headers = [
        "ФИО из табеля",
        "Смен",
        "Должная стадия",
        "Маркер",
        "Дата маркера",
        "Комментарий",
        "Таблица",
    ]
    worksheet.append(headers)
    for row in rows:
        sources = _source_urls(row.get("sources", []))
        worksheet.append(
            [
                row.get("employee"),
                row.get("shiftCount"),
                row.get("desiredStageId"),
                row.get("statusMarker"),
                row.get("statusMarkerDate"),
                row.get("statusComment"),
                "\n".join(sources),
            ]
        )
        _hyperlink_sources_cell(worksheet.cell(worksheet.max_row, 7), sources)
    _format_table(worksheet, widths=[38, 10, 18, 12, 14, 34, 82])
    workbook.save(path)


def _source_urls(sources: list[Any]) -> list[str]:
    urls = []
    for source in sources or []:
        sheet_id = str(source).split(":", 1)[0].strip()
        if sheet_id:
            urls.append(f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit")
    return sorted(set(urls))


def _hyperlink_cell(cell: Any) -> None:
    if cell.value:
        cell.hyperlink = str(cell.value)
        cell.style = "Hyperlink"


def _hyperlink_sources_cell(cell: Any, urls: list[str]) -> None:
    if len(urls) == 1:
        cell.hyperlink = urls[0]
        cell.style = "Hyperlink"
    elif urls:
        cell.font = Font(color="0563C1", underline="single")


def _format_table(worksheet: Any, *, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D9D9D9")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cells in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=worksheet.max_column,
    ):
        for cell in cells:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(1, worksheet.max_column).column_letter}{worksheet.max_row}"
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 34


if __name__ == "__main__":
    main()
