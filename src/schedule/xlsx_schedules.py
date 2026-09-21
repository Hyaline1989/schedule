from __future__ import annotations

from datetime import date
from itertools import chain
from pathlib import Path
from typing import BinaryIO
from warnings import warn

from openpyxl import load_workbook

from schedule.models import ParsedSchedule
from schedule.parsers.google_schedule import (
    MONTH_PATTERN,
    MONTH_ONLY_PATTERN,
    MONTHS,
    GoogleScheduleParser,
    ScheduleFormatError,
)


def parse_month_sheets(
    workbook_source: str | Path | BinaryIO,
    *,
    source_id: str,
    start_date: date | None = None,
    end_date: date | None = None,
) -> ParsedSchedule:
    workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    shifts = []
    settlements = []
    end_of_watches = []
    status_markers = []
    default_year = start_date.year if start_date else None
    parser = GoogleScheduleParser(default_year=default_year)
    try:
        explicit_periods = {
            period
            for worksheet in workbook.worksheets
            if (period := _worksheet_period(worksheet.title, default_year=None))
        }
        for worksheet in workbook.worksheets:
            period = _worksheet_period(
                worksheet.title,
                default_year=default_year,
                explicit_periods=explicit_periods,
            )
            if not period:
                continue
            sheet_month = date(period[0], period[1], 1)
            if start_date and sheet_month < start_date.replace(day=1):
                continue
            if end_date and sheet_month > end_date.replace(day=1):
                continue
            # Some real sheets keep the month and year only in the tab name.
            rows = chain(([worksheet.title],), worksheet.iter_rows(values_only=True))
            try:
                parsed = parser.parse_rows_data(
                    rows, source_id=f"{source_id}:{worksheet.title}"
                )
            except ScheduleFormatError as error:
                warn(
                    f"Лист {worksheet.title!r} пропущен: {error}",
                    RuntimeWarning,
                    stacklevel=2,
                )
                continue
            shifts.extend(
                item
                for item in parsed.shifts
                if (start_date is None or item.work_date >= start_date)
                and (end_date is None or item.work_date <= end_date)
            )
            settlements.extend(
                item
                for item in parsed.settlements
                if (start_date is None or item.settlement_date >= start_date)
                and (end_date is None or item.settlement_date <= end_date)
            )
            end_of_watches.extend(
                item
                for item in parsed.end_of_watches
                if (start_date is None or item.end_date >= start_date)
                and (end_date is None or item.end_date <= end_date)
            )
            status_markers.extend(
                item
                for item in parsed.status_markers
                if (start_date is None or item.marker_date >= start_date)
                and (end_date is None or item.marker_date <= end_date)
            )
    finally:
        workbook.close()
    return ParsedSchedule(
        tuple(shifts),
        tuple(settlements),
        tuple(end_of_watches),
        tuple(status_markers),
    )


def _worksheet_period(
    title: str,
    *,
    default_year: int | None = None,
    explicit_periods: set[tuple[int, int]] | None = None,
) -> tuple[int, int] | None:
    period = MONTH_PATTERN.search(title)
    if period:
        return int(period.group(2)), MONTHS[period.group(1).casefold()]
    if default_year:
        month = MONTH_ONLY_PATTERN.search(title)
        if month:
            candidate = (default_year, MONTHS[month.group(1).casefold()])
            if explicit_periods and candidate in explicit_periods:
                return None
            return candidate
    return None
