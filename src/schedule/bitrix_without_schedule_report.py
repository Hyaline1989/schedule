from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from schedule.assignment import OBJECT_FIELD
from schedule.bitrix import BitrixClient
from schedule.config import load_env_file
from schedule.models import normalize_employee_name
from schedule.name_aliases import candidate_names
from schedule.name_matching import find_smart_name_match
from schedule.sheet_objects import load_sheet_object_mappings
from schedule.sources import spreadsheet_ids
from schedule.xlsx_schedules import parse_month_sheets


START_DATE = date.fromisoformat(os.environ.get("SCHEDULE_START_DATE", "2026-08-01"))
END_DATE = date.fromisoformat(os.environ.get("SCHEDULE_END_DATE", date.today().isoformat()))

ACTIVE_STAGE_IDS = [
    "PREPARATION",  # Заселен
    "UC_L9233X",  # Отработал 1 смену
    "PREPAYMENT_INVOICE",  # Адаптирован
    "2",  # Больничный
]

STAGE_NAMES = {
    "3": "Направлен на объект",
    "PREPARATION": "Заселен",
    "UC_L9233X": "Отработал 1 смену",
    "PREPAYMENT_INVOICE": "Адаптирован",
    "2": "Больничный",
}


def main() -> None:
    load_env_file()
    runtime_dir = Path("runtime")
    output_dir = Path("outputs")
    output_dir.mkdir(exist_ok=True)

    mappings = load_sheet_object_mappings()
    company_by_sheet = {
        sheet_id: mapping.company_id for sheet_id, mapping in mappings.items()
    }
    sheet_by_company: dict[str, list[str]] = defaultdict(list)
    for sheet_id, mapping in mappings.items():
        sheet_by_company[mapping.company_id].append(sheet_id)

    schedule_names_by_company = _schedule_names_by_company(
        runtime_dir,
        company_by_sheet,
    )

    bitrix = BitrixClient.from_environment()
    deals = bitrix._list_deals({"CATEGORY_ID": 0, "@STAGE_ID": ACTIVE_STAGE_IDS})

    rows = []
    skipped: Counter[str] = Counter()
    checked_deals = 0
    known_company_ids = set(sheet_by_company)
    for deal in deals:
        if not _created_in_period(deal):
            skipped["created_outside_period"] += 1
            continue
        company_id = _object_id(deal)
        if company_id not in known_company_ids:
            skipped["object_not_in_google_scope"] += 1
            continue
        checked_deals += 1
        title = str(deal.get("TITLE") or "").strip()
        if _has_schedule_match(title, schedule_names_by_company.get(company_id, set())):
            skipped["found_in_schedule"] += 1
            continue
        rows.append(_row_payload(deal, company_id, sheet_by_company.get(company_id, [])))

    report = {
        "mode": "bitrix-without-schedule-report",
        "startDate": START_DATE.isoformat(),
        "endDate": END_DATE.isoformat(),
        "activeStageIds": ACTIVE_STAGE_IDS,
        "excludedStageIds": [
            "NEW",
            "3",
            "16",
            "LOSE",
            "WON",
            "UC_73ZTA3",
        ],
        "checkedDeals": checked_deals,
        "rowsCount": len(rows),
        "skipped": dict(skipped),
        "rows": rows,
    }
    report_path = runtime_dir / "bitrix-without-schedule-report.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    xlsx_path = output_dir / f"bitrix_without_schedule_{START_DATE.isoformat()}_{END_DATE.isoformat()}.xlsx"
    _write_xlsx(rows, xlsx_path)
    print(
        json.dumps(
            {
                "rows": len(rows),
                "report": str(report_path),
                "xlsx": str(xlsx_path),
                "byStage": dict(Counter(row["stage"] for row in rows).most_common()),
                "skipped": dict(skipped),
            },
            ensure_ascii=False,
        )
    )


def _schedule_names_by_company(
    runtime_dir: Path,
    company_by_sheet: dict[str, str],
) -> dict[str, set[str]]:
    result: dict[str, set[str]] = defaultdict(set)
    for spreadsheet_id in spreadsheet_ids():
        company_id = company_by_sheet.get(spreadsheet_id)
        if not company_id:
            continue
        path = runtime_dir / f"{spreadsheet_id}.xlsx"
        if not path.exists():
            continue
        parsed = parse_month_sheets(
            path,
            source_id=spreadsheet_id,
            start_date=START_DATE,
            end_date=END_DATE,
        )
        records = [
            *parsed.shifts,
            *parsed.settlements,
            *parsed.end_of_watches,
            *parsed.status_markers,
        ]
        for record in records:
            name = str(record.employee_name or "").strip()
            if name:
                result[company_id].add(name)
    return result


def _has_schedule_match(deal_title: str, schedule_names: set[str]) -> bool:
    normalized_candidates = {
        normalize_employee_name(candidate) for candidate in candidate_names(deal_title)
    }
    for schedule_name in schedule_names:
        if normalize_employee_name(schedule_name) in normalized_candidates:
            return True

    if find_smart_name_match(deal_title, sorted(schedule_names)):
        return True

    for schedule_name in sorted(schedule_names):
        reverse = find_smart_name_match(schedule_name, [deal_title])
        if reverse and normalize_employee_name(reverse.name) == normalize_employee_name(
            deal_title
        ):
            return True
    return False


def _created_in_period(deal: dict[str, Any]) -> bool:
    try:
        created = date.fromisoformat(str(deal.get("DATE_CREATE") or "")[:10])
    except ValueError:
        return False
    return START_DATE <= created <= END_DATE


def _object_id(deal: dict[str, Any]) -> str | None:
    value = deal.get(OBJECT_FIELD)
    if isinstance(value, list):
        value = value[0] if value else None
    if value is None:
        return None
    text = str(value)
    return text[3:] if text.startswith("CO_") else text


def _row_payload(
    deal: dict[str, Any],
    company_id: str,
    sheet_ids: list[str],
) -> dict[str, Any]:
    deal_id = str(deal.get("ID") or "")
    stage_id = str(deal.get("STAGE_ID") or "")
    return {
        "dealId": deal_id,
        "dealUrl": f"https://crm.complect.org/crm/deal/details/{deal_id}/",
        "employee": str(deal.get("TITLE") or "").strip(),
        "stageId": stage_id,
        "stage": STAGE_NAMES.get(stage_id, stage_id),
        "companyId": company_id,
        "companyUrl": f"https://crm.complect.org/crm/company/details/{company_id}/",
        "dateCreate": str(deal.get("DATE_CREATE") or "")[:10],
        "sheetUrls": "\n".join(
            f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit"
            for sheet_id in sheet_ids
        ),
    }


def _write_xlsx(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "no_schedule"
    worksheet.append(
        [
            "ФИО в Bitrix",
            "Ссылка на сделку",
            "Стадия",
            "ID объекта",
            "Ссылка на объект",
            "Дата создания",
            "Ссылка на таблицу",
        ]
    )
    for row in rows:
        worksheet.append(
            [
                row["employee"],
                row["dealUrl"],
                row["stage"],
                row["companyId"],
                row["companyUrl"],
                row["dateCreate"],
                row["sheetUrls"],
            ]
        )
        row_index = worksheet.max_row
        for column in (2, 5):
            cell = worksheet.cell(row_index, column)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
        table_cell = worksheet.cell(row_index, 7)
        urls = [item for item in str(row["sheetUrls"]).split("\n") if item]
        if len(urls) == 1:
            table_cell.hyperlink = urls[0]
            table_cell.style = "Hyperlink"
        elif urls:
            table_cell.font = Font(color="0563C1", underline="single")
        for column in range(1, 8):
            worksheet.cell(row_index, column).alignment = Alignment(
                vertical="top",
                wrap_text=column == 7,
            )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D9D9D9")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cells in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=7,
    ):
        for cell in cells:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:G{worksheet.max_row}"
    for index, width in enumerate([38, 48, 26, 12, 48, 14, 82], start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 32
    workbook.save(path)


if __name__ == "__main__":
    main()
