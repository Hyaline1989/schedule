from __future__ import annotations

import json
import os
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from schedule.bitrix import BitrixClient
from schedule.config import load_env_file
from schedule.sheet_objects import load_sheet_object_mappings


PUBLIC_REPORTS_DIR = Path(
    os.environ.get(
        "SCHEDULE_PUBLIC_REPORTS_DIR",
        "/var/www/maxbot_compl_usr/data/www/maxbot.complect.org/schedule-reports",
    )
)


def main() -> None:
    load_env_file()
    runtime_dir = Path("runtime")
    outputs_dir = Path("outputs")
    outputs_dir.mkdir(exist_ok=True)

    run_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    stage_report = _read_json(runtime_dir / "stage-mismatch-report.json")
    bitrix_without = _read_json(runtime_dir / "bitrix-without-schedule-report.json")
    apply_report = _latest_json(runtime_dir, "stage-mismatch-apply-*.json")

    stage_rows = list(stage_report.get("rows", []))
    apply_summary = apply_report.get("summary", {}) if apply_report else {}

    stage_path = outputs_dir / "stage_updates_summary.xlsx"
    missing_path = outputs_dir / "missing_by_object_summary.xlsx"

    _write_stage_updates(stage_path, stage_rows=stage_rows, apply_report=apply_report)
    _write_missing_by_object(
        missing_path,
        run_at=run_at,
        not_found_rows=[
            row
            for row in stage_report.get("unresolved", [])
            if row.get("reason") == "not_found"
        ],
        bitrix_without_rows=list(bitrix_without.get("rows", [])),
    )

    PUBLIC_REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(stage_path, PUBLIC_REPORTS_DIR / stage_path.name)
    shutil.copy2(missing_path, PUBLIC_REPORTS_DIR / missing_path.name)

    print(
        json.dumps(
            {
                "stageSummary": str(stage_path),
                "missingSummary": str(missing_path),
                "publicDir": str(PUBLIC_REPORTS_DIR),
            },
            ensure_ascii=False,
        )
    )


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _latest_json(directory: Path, pattern: str) -> dict[str, Any]:
    files = sorted(
        directory.glob(pattern),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    return _read_json(files[0]) if files else {}


def _write_stage_updates(path: Path, *, stage_rows: list[dict[str, Any]], apply_report: dict[str, Any]) -> None:
    rows_by_deal = {str(row.get("dealId") or ""): row for row in stage_rows}
    applied_rows = [
        row
        for row in apply_report.get("results", [])
        if row.get("updated") and str(row.get("dealId") or "") in rows_by_deal
    ]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Стадии"
    worksheet.append(["Ссылка на сделку", "Была стадия", "Стала стадия"])
    for applied in applied_rows:
        source_row = rows_by_deal[str(applied.get("dealId") or "")]
        deal_id = str(applied.get("dealId") or "")
        worksheet.append(
            [
                f"https://crm.complect.org/crm/deal/details/{deal_id}/",
                source_row.get("currentStage") or source_row.get("currentStageId"),
                source_row.get("desiredStage") or source_row.get("desiredStageId"),
            ]
        )
        cell = worksheet.cell(worksheet.max_row, 1)
        cell.hyperlink = str(cell.value)
        cell.style = "Hyperlink"
    _format_table(worksheet, widths=[52, 30, 30])
    workbook.save(path)


def _write_missing_by_object(
    path: Path,
    *,
    run_at: str,
    not_found_rows: list[dict[str, Any]],
    bitrix_without_rows: list[dict[str, Any]],
) -> None:
    object_names = _object_names()
    not_found_counts: Counter[str] = Counter()
    bitrix_without_counts: Counter[str] = Counter()

    for row in not_found_rows:
        company_id = _company_id_from_not_found(row)
        if company_id:
            not_found_counts[company_id] += 1

    for row in bitrix_without_rows:
        company_id = str(row.get("companyId") or "").strip()
        if company_id:
            bitrix_without_counts[company_id] += 1

    company_ids = sorted(
        set(not_found_counts) | set(bitrix_without_counts),
        key=lambda item: object_names.get(item, item),
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "По объектам"
    worksheet.append(
        ["Дата/время", "Объект", "Не хватает в Битрикс", "Не хватает в табеле"]
    )

    for company_id in company_ids:
        worksheet.append(
            [
                run_at,
                object_names.get(company_id, f"ID {company_id}"),
                not_found_counts.get(company_id, 0),
                bitrix_without_counts.get(company_id, 0),
            ]
        )

    _format_table(worksheet, widths=[22, 36, 22, 22])
    workbook.save(path)


def _company_id_from_not_found(row: dict[str, Any]) -> str:
    mappings = load_sheet_object_mappings()
    for source in row.get("sources") or []:
        sheet_id = str(source).split(":", 1)[0]
        mapping = mappings.get(sheet_id)
        if mapping:
            return str(mapping.company_id)
    return ""


def _object_names() -> dict[str, str]:
    result = {}
    for mapping in load_sheet_object_mappings().values():
        result[str(mapping.company_id)] = mapping.object_name or f"ID {mapping.company_id}"
    try:
        bitrix = BitrixClient.from_environment()
    except Exception:
        return result
    for company_id in sorted(result):
        try:
            company = bitrix.get_company(company_id)
        except Exception:
            continue
        title = str(company.get("TITLE") or "").strip()
        if title:
            result[company_id] = title
    return result


def _format_table(worksheet: Any, *, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D9D9D9")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cells in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=worksheet.max_column,
    ):
        for cell in cells:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{worksheet.cell(1, worksheet.max_column).column_letter}{worksheet.max_row}"
    )
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width


if __name__ == "__main__":
    main()
