from __future__ import annotations

import json
import os
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PORTAL_URL = os.environ.get("BITRIX_PORTAL_URL", "https://crm.complect.org").rstrip("/")


def main() -> None:
    start_date = date.fromisoformat(os.environ.get("SCHEDULE_START_DATE", "2026-08-01"))
    end_date = date.fromisoformat(os.environ.get("SCHEDULE_END_DATE", date.today().isoformat()))
    runtime_dir = Path("runtime")
    output_dir = Path("outputs")
    output_dir.mkdir(exist_ok=True)

    issues = json.loads(
        (runtime_dir / "new-notification-issues-2026-09-21-preview.json").read_text(
            encoding="utf-8"
        )
    )
    rows = [row for row in issues.get("rows", []) if row.get("isNew")]
    bitrix_without_schedule = [
        row for row in rows if row.get("eventType") == "bitrix_without_schedule"
    ]
    not_found_in_bitrix = [
        row for row in rows if row.get("eventType") == "not_found_in_bitrix"
    ]

    bitrix_path = (
        output_dir
        / f"bitrix_without_schedule_new_only_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"
    )
    not_found_path = (
        output_dir
        / f"not_found_in_bitrix_new_only_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"
    )
    _write_bitrix_without_schedule(bitrix_without_schedule, bitrix_path)
    _write_not_found_in_bitrix(not_found_in_bitrix, not_found_path)
    print(
        json.dumps(
            {
                "bitrixWithoutScheduleNew": len(bitrix_without_schedule),
                "bitrixWithoutScheduleXlsx": str(bitrix_path),
                "notFoundInBitrixNew": len(not_found_in_bitrix),
                "notFoundInBitrixXlsx": str(not_found_path),
            },
            ensure_ascii=False,
        )
    )


def _write_bitrix_without_schedule(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "new_only"
    headers = [
        "ФИО в Bitrix",
        "Сделка",
        "Стадия",
        "ID объекта",
        "Дата создания",
        "Таблица",
    ]
    worksheet.append(headers)
    for row in rows:
        state = row.get("state") or {}
        deal_id = str(row.get("dealId") or "")
        worksheet.append(
            [
                row.get("employee"),
                f"{PORTAL_URL}/crm/deal/details/{deal_id}/" if deal_id else "",
                state.get("stage"),
                row.get("objectKey"),
                state.get("dateCreate"),
                state.get("sheetUrls"),
            ]
        )
        _hyperlink_cell(worksheet.cell(worksheet.max_row, 2))
        _hyperlink_sources_cell(worksheet.cell(worksheet.max_row, 6))
    _format_table(worksheet, widths=[38, 48, 24, 12, 14, 82])
    workbook.save(path)


def _write_not_found_in_bitrix(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "new_only"
    headers = [
        "ФИО из табеля",
        "ID объекта",
        "Смен",
        "Должная стадия",
        "Маркер",
        "Дата маркера",
        "Таблица",
    ]
    worksheet.append(headers)
    for row in rows:
        state = row.get("state") or {}
        worksheet.append(
            [
                row.get("employee"),
                row.get("objectKey"),
                state.get("shiftCount"),
                state.get("desiredStageId"),
                state.get("statusMarker"),
                state.get("statusMarkerDate"),
                "\n".join(_source_urls(state.get("sources") or [])),
            ]
        )
        _hyperlink_sources_cell(worksheet.cell(worksheet.max_row, 7))
    _format_table(worksheet, widths=[38, 12, 10, 18, 12, 14, 82])
    workbook.save(path)


def _source_urls(sources: list[Any]) -> list[str]:
    urls = []
    for source in sources:
        sheet_id = str(source).split(":", 1)[0].strip()
        if sheet_id:
            urls.append(f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit")
    return sorted(set(urls))


def _hyperlink_cell(cell: Any) -> None:
    if cell.value:
        cell.hyperlink = str(cell.value)
        cell.style = "Hyperlink"


def _hyperlink_sources_cell(cell: Any) -> None:
    urls = [item for item in str(cell.value or "").split("\n") if item]
    if len(urls) == 1:
        cell.hyperlink = urls[0]
        cell.style = "Hyperlink"
    elif urls:
        cell.font = Font(color="0563C1", underline="single")


def _format_table(worksheet: Any, *, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D9D9D9")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cells in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=worksheet.max_column,
    ):
        for cell in cells:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{worksheet.cell(1, worksheet.max_column).column_letter}{worksheet.max_row}"
    )
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 34


if __name__ == "__main__":
    main()
